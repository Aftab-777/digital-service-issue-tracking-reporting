"""Recheck the fixed 7 October 2026 portfolio snapshot.

Run from anywhere after installing validation/requirements.txt:
    python validation/check_project.py
"""

from __future__ import annotations

import csv
import json
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
ASOF = date(2026, 10, 7)


def check(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def issue(row: dict[str, str]) -> bool:
    received = date.fromisoformat(row["Date received"])
    target = date.fromisoformat(row["Target completion date"]) if row["Target completion date"] else None
    updated = date.fromisoformat(row["Last update date"]) if row["Last update date"] else None
    closed = date.fromisoformat(row["Closed date"]) if row["Closed date"] else None
    return (target is None or updated is None or target < received or updated < received or updated > ASOF
            or (row["Status"] == "Completed" and (closed is None or closed < received or closed > ASOF))
            or (row["Status"] != "Completed" and closed is not None))


with (ROOT / "data/requests.csv").open(encoding="utf-8-sig", newline="") as stream:
    rows = list(csv.DictReader(stream))
metrics = json.loads((ROOT / "data/metrics.json").read_text(encoding="utf-8"))
ids = [row["Request ID"] for row in rows]
check(len(rows) == 36 and len(set(ids)) == 36, "Expected 36 unique requests")
check(metrics["as_of"] == ASOF.isoformat(), "Reporting date changed")
check(metrics["total"] == len(rows), "Total does not match CSV")
for field, column in [("status", "Status"), ("category", "Request category"), ("priority", "Priority")]:
    check(metrics[field] == dict(Counter(row[column] for row in rows)), f"{field} counts differ")

open_rows = [row for row in rows if row["Status"] != "Completed"]
overdue_ids = [row["Request ID"] for row in open_rows if row["Target completion date"] and date.fromisoformat(row["Target completion date"]) < ASOF]
missing_ids = [row["Request ID"] for row in rows if not row["Assigned owner"]]
issue_ids = [row["Request ID"] for row in rows if issue(row)]
expected = {
    "open": len(open_rows), "overdue": len(overdue_ids), "blocked": sum(row["Status"] == "Blocked" for row in rows),
    "missing_owner": len(missing_ids), "data_issue": len(issue_ids),
    "high_open": sum(row["Priority"] == "High" for row in open_rows),
}
for key, value in expected.items():
    check(metrics[key] == value, f"{key} differs from CSV")
for key, value in [("overdue_ids", overdue_ids), ("missing_owner_ids", missing_ids), ("data_issue_ids", issue_ids)]:
    check(metrics[key] == value, f"{key} list differs")
check(missing_ids == ["DS-021", "DS-032"], "Missing-owner examples changed")
check(issue_ids == ["DS-033", "DS-034"], "Data-quality examples changed")
check("DS-034" in overdue_ids, "Provisional overdue explanation changed")

book = ROOT / "workbook/Digital_Service_Issue_Tracker.xlsx"
formula_book = load_workbook(book, data_only=False)
values_book = load_workbook(book, data_only=True)
check(formula_book.sheetnames == ["Summary", "Requests", "Instructions"], "Workbook sheets changed")
check("RequestsTable" in formula_book["Requests"].tables, "Request table missing")
check(len(formula_book["Requests"].data_validations.dataValidation) == 2, "Dropdowns changed")
check(formula_book["Instructions"]["A1"].value == "Operations request tracker", "Workbook title changed")
reported_date = values_book["Instructions"]["B4"].value
check(isinstance(reported_date, (date, datetime)) and reported_date.date() == ASOF if isinstance(reported_date, datetime) else reported_date == ASOF, "Workbook as-of date changed")
for column in "NOPQR":
    check(formula_book["Requests"][f"{column}7"].value.startswith("=IF("), f"{column} formula missing")
summary = values_book["Summary"]
for cell, key in [("B6", "total"), ("D6", "open"), ("F6", "overdue"), ("H6", "blocked"),
                  ("B8", "missing_owner"), ("D8", "data_issue"), ("F8", "high_open")]:
    check(summary[cell].value == metrics[key], f"Summary {cell} differs")
check(summary["H8"].value == metrics["status"]["Completed"], "Completed count differs")
check(summary["B19"].value == "PASS", "Status reconciliation failed")
check([summary[f"B{i}"].value for i in range(12, 16)] == [metrics["status"][s] for s in ["New", "In Progress", "Blocked", "Completed"]], "Status breakdown differs")
check(values_book["Requests"]["P27"].value == "Yes" and values_book["Requests"]["P38"].value == "Yes", "Missing-owner flags differ")
check(values_book["Requests"]["Q39"].value == "Review" and values_book["Requests"]["Q40"].value == "Review", "Data-issue flags differ")

report_texts = []
for stem, title in [("Status_Report", "Weekly Request Status Report"), ("Decision_Briefing", "Request Prioritization Brief")]:
    pdf = PdfReader(str(ROOT / "reports" / f"{stem}.pdf"))
    check(len(pdf.pages) == 1, f"{stem} is not one page")
    text = pdf.pages[0].extract_text()
    report_texts.append(text)
    check(title in text and "synthetic" in text.lower(), f"{stem} title/disclosure missing")
    for key in ["total", "open", "overdue", "blocked"]:
        check(str(metrics[key]) in text, f"{stem} missing {key}")
    check(title in (ROOT / "reports" / f"{stem}.md").read_text(encoding="utf-8"), f"{stem} Markdown differs")

with zipfile.ZipFile(ROOT / "presentation/Management_Briefing.pptx") as archive:
    slide_names = sorted((name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
                         key=lambda name: int(re.search(r"\d+", name).group()))
    check(len(slide_names) == 5, "Expected five slides")
    slides = []
    for name in slide_names:
        xml = ET.fromstring(archive.read(name))
        slides.append(" ".join(node.text or "" for node in xml.iter() if node.tag.endswith("}t")))
    notes = " ".join(archive.read(name).decode("utf-8") for name in archive.namelist() if name.startswith("ppt/notesSlides/") and name.endswith(".xml"))
check("Operations Request Tracker and Reporting" in slides[0], "Slide title differs")
check(all("Personal portfolio simulation" in slide for slide in slides), "Slide disclosure missing")
check(all(str(metrics[key]) in slides[1] for key in ["total", "open", "overdue", "blocked"]), "Slide figures differ")
check("DS-033" in slides[2] and "DS-034" in slides[2], "Quality flags missing from slides")
check("option 2" in slides[3].lower(), "Recommendation missing from slides")

readme = (ROOT / "README.md").read_text(encoding="utf-8")
headings = re.findall(r"^## (.+)$", readme, flags=re.MULTILINE)
check(headings == ["Project overview", "Business problem", "Intended users", "Project scope", "How the tracker works",
                   "Reporting outputs", "Main findings", "Design decisions and limitations", "How to use the files", "Future improvements"], "README structure differs")
for target in re.findall(r"!?(?:\[[^\]]+\])\(([^)]+)\)", readme):
    if not target.startswith(("http://", "https://", "#")):
        check((ROOT / target).is_file(), f"README link missing: {target}")
for name in ["tracker-summary", "tracker-requests", "tracker-instructions", "manager-options", "status-report"]:
    image = ROOT / "screenshots" / f"{name}.png"
    check(image.read_bytes().startswith(b"\x89PNG\r\n\x1a\n"), f"Capture missing: {name}")

public_text = "\n".join(path.read_text(encoding="utf-8-sig") for path in ROOT.rglob("*")
                        if path.is_file() and path.suffix in {".md", ".json", ".csv"})
public_text += "\n" + "\n".join(report_texts + slides + [notes])
for package in [book, ROOT / "presentation/Management_Briefing.pptx"]:
    with zipfile.ZipFile(package) as archive:
        public_text += "\n" + "\n".join(archive.read(name).decode("utf-8", errors="ignore")
                                          for name in archive.namelist() if name.endswith(".xml"))
        themes = [name for name in archive.namelist() if "/theme" in name and name.endswith(".xml")]
        check(bool(themes), f"Theme metadata missing from {package.name}")
        for name in themes:
            theme = ET.fromstring(archive.read(name))
            check(theme.attrib.get("name") == "Operations Request Tracker", f"Theme name changed in {name}")
            for element in theme.iter():
                if element.tag.endswith(("}clrScheme", "}fmtScheme")):
                    check(element.attrib.get("name") == "Operations Request Tracker", f"Scheme name changed in {name}")
        if package.suffix == ".pptx":
            core = ET.fromstring(archive.read("docProps/core.xml"))
            app = ET.fromstring(archive.read("docProps/app.xml"))
            check(not any(node.tag.endswith(("}creator", "}lastModifiedBy")) for node in core),
                  "Presentation creator property changed")
            check(not any(node.tag.endswith("}Application") for node in app),
                  "Presentation application property changed")
for pattern in [r"\bgovernment\s+of\s+[a-z]+\b", r"\brequisition\s+\d+\b",
                r"\bjob\s+(?:post\w+|description)\b", r"\bapplication\s+deadline\b"]:
    check(re.search(pattern, public_text, flags=re.IGNORECASE) is None, f"Unwanted text or package metadata remains: {pattern}")

print("PASS: 36 synthetic records, workbook figures/formulas, reports, slides, captures, links and package metadata.")
